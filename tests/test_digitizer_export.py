"""Exportación autónoma X–Y: coherencia exacta CSV ↔ XLSX ↔ JSON, alcance, paneles y unidades."""
import csv
import io
import json
import zipfile
from pathlib import Path

import numpy as np
import pytest
from openpyxl import load_workbook
from PIL import Image

from digitizer.detection import DetectionParams
from digitizer.draft import DigitizationDraft
from digitizer.export import (
    CURRENT, ExportError, collect, entries_for, to_csv, to_json, to_png, to_xlsx, to_zip,
)
from digitizer.imageio import load_image

ROOT = Path(__file__).resolve().parents[1]
FX = ROOT / "tests" / "fixtures" / "digitizer"


def _two_curves():
    d = DigitizationDraft.from_image(load_image((FX / "02_dos_curvas.png").read_bytes(), "02_dos_curvas.png"))
    lines = [c for c in d.detection.candidates if c.kind == "línea"]
    for name, c in zip(("A", "B"), lines):
        d.select_candidate(c.id)
        d.save_curve(name)
    return d


def _parse_all(b):
    rows_csv = list(csv.DictReader(io.StringIO(to_csv(b).decode("utf-8"))))
    wb = load_workbook(io.BytesIO(to_xlsx(b)))
    sh = wb["Datos"]
    head = [c.value for c in sh[1]]
    rows_x = [dict(zip(head, [c.value for c in r])) for r in sh.iter_rows(min_row=2)]
    js = json.loads(to_json(b))
    rows_j = [dict(p, curve_id=c["curve_id"]) for pnl in js["panels"] for c in pnl["curves"] for p in c["points"]]
    return rows_csv, rows_x, rows_j, wb, js


def test_csv_xlsx_json_are_identical_value_by_value():
    d = _two_curves()
    b = collect(entries_for([d], "panel", d))
    rows_csv, rows_x, rows_j, wb, js = _parse_all(b)
    assert len(rows_csv) == len(rows_x) == len(rows_j) == len(b.rows) == 80
    for rc, rx, rj, r in zip(rows_csv, rows_x, rows_j, b.rows):
        for k in ("x", "y", "px", "py"):
            assert float(rc[k]) == rx[k] == rj[k] == r[k]                 # exactitud bit a bit
        assert rc["curve_id"] == rx["curve_id"] == rj["curve_id"] == r["curve_id"]
        assert rc["x_unit"] == rx["x_unit"] == "A/cm²" and rc["y_unit"] == rx["y_unit"] == "V"
    assert {r["curve_id"] for r in b.rows} == {"P1_A", "P1_B"}
    assert {"Datos", "Metadatos", "P1_A", "P1_B"} <= set(wb.sheetnames)
    assert wb["P1_A"]["B1"].value == "x [A/cm²]"
    # los datos exportados son exactamente los del estado del borrador (calibración vigente)
    cal = d.calibration()[0]
    pts = sorted([p for p in d.saved_curves[0]["points"] if p["include"]], key=lambda p: (p["px"], p["py"]))
    xs, ys = cal.to_physical([p["px"] for p in pts], [p["py"] for p in pts])
    a_rows = [r for r in b.rows if r["curve_id"] == "P1_A"]
    assert [r["x"] for r in a_rows] == [float(v) for v in xs] and [r["y"] for r in a_rows] == [float(v) for v in ys]
    meta = js["panels"][0]
    assert meta["axis_confirmed"] == {"X": True, "Y": True} and meta["units"]["confirmed"]
    assert meta["axis_recognition"]["X"]["status"] == "CONFIABLE"


def test_scopes_current_trace_single_saved_and_errors():
    d = _two_curves()
    b = collect(entries_for([d], "curva", d, CURRENT))
    assert {r["curve_id"] for r in b.rows} == {"P1_curva"}
    b = collect(entries_for([d], "curva", d, "B"))
    assert {r["curve_id"] for r in b.rows} == {"P1_B"}
    with pytest.raises(ExportError, match="No existe"):
        entries_for([d], "curva", d, "Z")
    empty = DigitizationDraft.from_image(load_image((FX / "01_curva_simple.png").read_bytes(), "x.png"))
    with pytest.raises(ExportError, match="Guardar curva"):
        entries_for([empty], "panel", empty)


def test_excluded_points_are_not_exported_but_stay_in_audit():
    d = DigitizationDraft.from_image(load_image((FX / "01_curva_simple.png").read_bytes(), "x.png"))
    d.apply_event({"type": "toggle_include", "id": d.points[0].id})
    d.save_curve("A")
    b = collect(entries_for([d], "panel", d))
    assert len(b.rows) == 39
    audit = json.loads(to_json(b))
    assert audit["panels"][0]["curves"][0]["name"] == "A"


def test_unconfirmed_axis_blocks_physical_export_pixels_only_is_explicit():
    d = DigitizationDraft.from_image(load_image((FX / "01_curva_simple.png").read_bytes(), "x.png"), ocr=False)
    d.save_curve("A")
    with pytest.raises(ExportError, match="no se pueden exportar valores físicos"):
        collect(entries_for([d], "panel", d))
    b = collect(entries_for([d], "panel", d), pixels_only=True)
    rows = list(csv.DictReader(io.StringIO(to_csv(b).decode())))
    assert set(rows[0]) == {"curve_id", "panel", "point", "px", "py", "origin", "status"}
    assert rows[0]["status"] == "SIN CALIBRAR"                                  # ni x, ni y, ni unidades
    assert json.loads(to_json(b))["pixels_only"] is True


def test_free_units_and_png_zip_contents():
    d = DigitizationDraft.from_image(load_image((FX / "01_curva_simple.png").read_bytes(), "01_curva_simple.png"))
    d.set_free_units("mA/cm²", "mV")          # el usuario reetiqueta (la escala la dan las anclas, no la unidad)
    d.save_curve("A")
    b = collect(entries_for([d], "panel", d))
    assert {(r["x_unit"], r["y_unit"]) for r in b.rows} == {("mA/cm²", "mV")}
    Image.open(io.BytesIO(to_png(b))).verify()
    z = zipfile.ZipFile(io.BytesIO(to_zip(b)))
    names = set(z.namelist())
    assert {"LEEME.txt", "datos.csv", "datos.xlsx", "datos.json", "overlay.png", "original/01_curva_simple.png",
            "auditoria/panel_1.json"} <= names
    assert z.read("datos.csv") == to_csv(b)
    assert z.read("original/01_curva_simple.png") == d.image.original_bytes


def test_real_multipanel_figure_each_panel_keeps_its_own_calibration():
    img = load_image((ROOT / "examples" / "digitizer" / "otekon_figure3_source.png").read_bytes(), "otekon.png")
    p1 = DigitizationDraft.from_image(img, DetectionParams(panel_index=0))
    p4 = DigitizationDraft.from_image(img, DetectionParams(panel_index=3))
    for d in (p1, p4):
        assert d.physical_ready()[0], d.physical_ready()[1]
        blue = next(c for c in d.detection.candidates if c.color_name == "azul" and c.confidence >= 0.5)
        d.select_candidate(blue.id)
        d.save_curve("azul")
    b = collect(entries_for([p1, p4], "imagen", p1))
    js = json.loads(to_json(b))
    assert [p["panel"] for p in js["panels"]] == [1, 4]
    c1, c4 = js["panels"][0]["calibration"], js["panels"][1]["calibration"]
    assert c1["x"]["pixel_1"] != c4["x"]["pixel_1"]                            # ejes de paneles distintos
    xs4 = [r["x"] for r in b.rows if r["panel"] == 4]
    assert 0.0 <= min(xs4) and max(xs4) <= 1.2
    assert {r["curve_id"] for r in b.rows} == {"P1_azul", "P4_azul"}
