import asyncio, json, io, zipfile, csv
from playwright.async_api import async_playwright
import os, pathlib
U = os.environ.get("PML_URL", "http://127.0.0.1:8501"); S = str(pathlib.Path(__file__).resolve().parents[2] / "docs" / "capturas" / "v0.3.0") + "/"; os.makedirs(S, exist_ok=True)
R={}
async def idle(pg, ms=900):
    await pg.wait_for_timeout(ms)
    for _ in range(160):
        if await pg.locator('[data-testid="stStatusWidget"]').count()==0: break
        await pg.wait_for_timeout(250)
    await pg.wait_for_timeout(500)
async def txt(pg): return await pg.inner_text("body")
async def main():
    async with async_playwright() as p:
        b = await p.chromium.launch(args=["--no-proxy-server"]); ctx = await b.new_context(viewport={"width":1440,"height":1000}, accept_downloads=True)
        pg = await ctx.new_page(); errs=[]; pg.on("pageerror", lambda e: errs.append(str(e)))
        try:
            await pg.goto(U + "/?view=digitizer"); await idle(pg, 3500)
            await pg.screenshot(path=S+"20_digitalizador_inicio.png")
            await pg.get_by_role("button", name="Figura publicada OTEKON Fig. 3").click(); await idle(pg, 6000)
            t = await txt(pg)
            R["eje_X"] = [l for l in t.splitlines() if l.startswith("Eje X ·")][:1]
            R["eje_Y"] = [l for l in t.splitlines() if l.startswith("Eje Y ·")][:1]
            R["calibrado"] = "Calibrado y confirmado" in t
            R["unidades"] = [l for l in t.splitlines() if "Unidades confirmadas" in l][:1]
            await pg.screenshot(path=S+"21_digitalizador_reconocimiento.png")
            await pg.get_by_role("button", name="Guardar curva").click(); await idle(pg, 2000)
            # panel 4
            await pg.get_by_label("Panel", exact=True).click(); await idle(pg,400)
            await pg.get_by_role("option", name="Panel 4").click(); await idle(pg, 6000)
            t = await txt(pg); R["panel4_calibrado"] = "Calibrado y confirmado" in t
            await pg.screenshot(path=S+"22_digitalizador_panel4.png")
            await pg.get_by_role("button", name="Guardar curva").click(); await idle(pg, 2000)
            await pg.get_by_label("Panel", exact=True).click(); await idle(pg,400)
            await pg.get_by_role("option", name="Panel 1 · trabajado").click(); await idle(pg, 3000)
            t = await txt(pg); R["vuelta_panel1_curvas"] = [l for l in t.splitlines() if "curva_1 ·" in l][:1]
            await pg.get_by_role("tab", name="Exportar datos").click(); await idle(pg)
            await pg.get_by_text("Todas las curvas guardadas de la imagen").click(); await idle(pg, 2000)
            await pg.screenshot(path=S+"23_digitalizador_exportar.png")
            await pg.get_by_role("button", name="Preparar archivos de exportación").click(); await idle(pg, 4000)
            files = {}
            for k in ("CSV", "XLSX", "JSON", "PNG", "ZIP"):
                async with pg.expect_download() as dl: await pg.get_by_role("button", name=k, exact=True).click()
                d = await dl.value; files[k] = open(await d.path(), "rb").read(); R[f"dl_{k}"] = (d.suggested_filename, len(files[k]))
                await idle(pg, 600)
            rows = list(csv.DictReader(io.StringIO(files["CSV"].decode())))
            js = json.loads(files["JSON"])
            R["csv_curvas"] = sorted({r["curve_id"] for r in rows}); R["csv_filas"] = len(rows)
            R["json_paneles"] = [p["panel"] for p in js["panels"]]
            jrows = [pt for pn in js["panels"] for c in pn["curves"] for pt in c["points"]]
            R["csv_igual_json"] = all(float(a["x"]) == b_["x"] and float(a["y"]) == b_["y"] for a, b_ in zip(rows, jrows)) and len(rows) == len(jrows)
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(files["XLSX"])); sh = wb["Datos"]; head=[c.value for c in sh[1]]
            xr = [dict(zip(head,[c.value for c in r])) for r in sh.iter_rows(min_row=2)]
            R["csv_igual_xlsx"] = all(float(a["x"]) == x["x"] and float(a["y"]) == x["y"] for a, x in zip(rows, xr)) and len(rows) == len(xr)
            R["zip"] = sorted(zipfile.ZipFile(io.BytesIO(files["ZIP"])).namelist())
            R["unidades_csv"] = sorted({(r["x_unit"], r["y_unit"]) for r in rows})
        finally:
            R["exc"] = await pg.locator('[data-testid="stException"]').count(); R["errs"] = errs[:4]
            print(json.dumps(R, indent=1, ensure_ascii=False)); await b.close()
asyncio.run(main())
