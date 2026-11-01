"""Exportación autónoma de datos X–Y digitalizados (sin Motor 1 ni perfil PEMFC).

Una única recolección (``collect``) produce la tabla de puntos; CSV, XLSX, JSON, PNG y ZIP se generan
a partir de ella, de modo que todos los formatos contienen EXACTAMENTE los mismos valores (floats en
representación ``repr`` → ida y vuelta bit a bit).

Reglas: solo se exportan valores físicos de paneles con calibración confirmada (ejes y unidades). La
exportación en píxeles es explícita y queda marcada «SIN CALIBRAR»; nunca se inventan unidades.
"""
from __future__ import annotations

import csv
import io
import json
import math
import re
import zipfile
from dataclasses import dataclass, field
from datetime import datetime, timezone

import numpy as np
from PIL import Image, ImageDraw

from pemfc_core.schemas import APP_VERSION

from . import ALGORITHM_ID, DIGITIZER_VERSION
from .draft import DigitizationDraft, _jsonable

SCHEMA = "pemfc-modellab/xy-export"
SCHEMA_VERSION = "1.0.0"
COLUMNS = ["curve_id", "panel", "point", "x", "y", "x_unit", "y_unit", "origin", "px", "py"]
PX_COLUMNS = ["curve_id", "panel", "point", "px", "py", "origin", "status"]
PALETTE = [(0, 150, 140), (47, 111, 237), (224, 137, 30), (122, 79, 209), (209, 73, 91), (60, 157, 93)]
CURRENT = "(actual)"


class ExportError(ValueError):
    pass


@dataclass
class CurveEntry:
    curve_id: str
    panel: int
    name: str
    draft: DigitizationDraft
    points: list[dict]                   # DraftPoint.to_dict() incluidos, ordenados por px
    candidate: str | None = None
    method: str = ""


@dataclass
class Bundle:
    entries: list[CurveEntry]
    pixels_only: bool
    rows: list[dict] = field(default_factory=list)
    created_utc: str = field(default_factory=lambda: datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"))

    @property
    def image(self):
        return self.entries[0].draft.image


def _units(d: DigitizationDraft) -> tuple[str, str]:
    return (d.x_unit_text or d.current_unit, d.y_unit_text or d.voltage_unit)


def _panel_label(d: DigitizationDraft) -> int:
    return int(d.panel_index + 1) if d.panel_index is not None else 1


def entries_for(drafts: list[DigitizationDraft], scope: str, active: DigitizationDraft,
                curve_name: str | None = None) -> list[CurveEntry]:
    """scope: «curva» (la indicada o la traza actual), «panel» (guardadas del panel activo) o «imagen»."""
    def mk(d, name, pts, cand, method):
        pts = sorted([p for p in pts if p.get("include", True)], key=lambda p: (p["px"], p["py"]))
        safe = re.sub(r"[^\w.\-]+", "_", name).strip("_") or "curva"
        return CurveEntry(f"P{_panel_label(d)}_{safe}", _panel_label(d), name, d, pts, cand, method)

    out = []
    if scope == "curva":
        if curve_name in (None, CURRENT):
            out.append(mk(active, "curva", [p.to_dict() for p in active.points], active.selected_candidate,
                          active.extraction_method))
        else:
            c = next((c for c in active.saved_curves if c["name"] == curve_name), None)
            if c is None:
                raise ExportError(f"No existe la curva {curve_name!r} en el panel activo.")
            out.append(mk(active, c["name"], c["points"], c.get("candidate"), c.get("method", "")))
    else:
        for d in ([active] if scope == "panel" else drafts):
            for c in d.saved_curves:
                out.append(mk(d, c["name"], c["points"], c.get("candidate"), c.get("method", "")))
    if not out:
        raise ExportError("No hay curvas guardadas que exportar en este alcance. Use «Guardar curva» primero.")
    ids = [e.curve_id for e in out]
    if len(set(ids)) != len(ids):
        raise ExportError("Identificadores de curva duplicados en la exportación.")
    if any(not e.points for e in out):
        raise ExportError("Alguna curva no tiene puntos incluidos.")
    return out


def collect(entries: list[CurveEntry], pixels_only: bool = False) -> Bundle:
    b = Bundle(entries, pixels_only)
    for e in entries:
        d = e.draft
        px = np.array([p["px"] for p in e.points], float)
        py = np.array([p["py"] for p in e.points], float)
        if pixels_only:
            for i, p in enumerate(e.points):
                b.rows.append({"curve_id": e.curve_id, "panel": e.panel, "point": i + 1, "px": float(px[i]),
                               "py": float(py[i]), "origin": "manual" if p["method"] == "manual" or p["edits"] else "automático",
                               "status": "SIN CALIBRAR"})
            continue
        ok, why = d.physical_ready()
        if not ok:
            raise ExportError(f"Panel {e.panel}: no se pueden exportar valores físicos: " + " ".join(why))
        cal = d.calibration()[0]
        xs, ys = cal.to_physical(px, py)
        xu, yu = _units(d)
        for i, p in enumerate(e.points):
            b.rows.append({"curve_id": e.curve_id, "panel": e.panel, "point": i + 1, "x": float(xs[i]),
                           "y": float(ys[i]), "x_unit": xu, "y_unit": yu,
                           "origin": "manual" if p["method"] == "manual" or p["edits"] else "automático",
                           "px": float(px[i]), "py": float(py[i])})
    return b


# --------------------------------------------------------------------------- formatos
def to_csv(b: Bundle) -> bytes:
    buf = io.StringIO()
    cols = PX_COLUMNS if b.pixels_only else COLUMNS
    w = csv.DictWriter(buf, fieldnames=cols, lineterminator="\n")
    w.writeheader()
    for r in b.rows:
        w.writerow({k: (repr(v) if isinstance(v, float) else v) for k, v in r.items() if k in cols})
    return buf.getvalue().encode("utf-8")


def _panel_meta(d: DigitizationDraft) -> dict:
    cal, errs, warns = d.calibration()
    xu, yu = _units(d)
    return {
        "panel": _panel_label(d), "plot_box_px": list(d.plot_box) if d.plot_box else None,
        "roi_px": list(d.roi) if d.roi else None,
        "calibration": cal.to_dict() if cal else None, "calibration_errors": errs,
        "anchors": [d.anchors[n].to_dict() for n in ("X1", "X2", "Y1", "Y2")],
        "axis_scales": dict(d.scales), "axis_confirmed": dict(d.axis_confirmed),
        "axis_recognition": {k: {"status": v.status, "reasons": v.reasons, "scale": v.scale, "n_inliers": v.n_inliers,
                                 "rms_px": v.rms_px, "multiplier": v.multiplier, "engine": v.engine,
                                 "labels": [{"text": l.text, "value": l.effective_value, "inlier": l.inlier,
                                             "bbox_px": l.bbox_px, "user_value": l.user_value,
                                             "excluded": l.excluded_by_user} for l in v.labels],
                                 "title": v.title} for k, v in d.axis_recognition.items()},
        "units": {"x": xu, "y": yu, "confirmed": d.units_confirmed, "source": d.units_source},
    }


def to_json(b: Bundle) -> bytes:
    img = b.image
    panels: dict[int, dict] = {}
    for e in b.entries:
        p = panels.setdefault(e.panel, {**_panel_meta(e.draft), "curves": []})
        rows = [r for r in b.rows if r["curve_id"] == e.curve_id]
        p["curves"].append({"curve_id": e.curve_id, "name": e.name, "candidate": e.candidate, "method": e.method,
                            "points": [{k: v for k, v in r.items() if k not in ("curve_id", "panel")} for r in rows]})
    doc = {"schema": SCHEMA, "schema_version": SCHEMA_VERSION, "app_version": APP_VERSION,
           "engine": {"algorithm": ALGORITHM_ID, "version": DIGITIZER_VERSION}, "created_utc": b.created_utc,
           "pixels_only": b.pixels_only,
           "pixel_convention": "origen arriba-izquierda; y hacia ABAJO; coordenadas continuas",
           "image": img.describe(), "panels": [panels[k] for k in sorted(panels)],
           "notice": ("Datos DIGITALIZADOS de una imagen: no constituyen una medición independiente. "
                      + ("SIN CALIBRAR: solo coordenadas en píxeles." if b.pixels_only else
                         "Valores físicos derivados de las referencias de eje confirmadas."))}
    return json.dumps(_jsonable(doc), ensure_ascii=False, indent=1).encode("utf-8")


def _sheet_title(name: str, used: set) -> str:
    base = re.sub(r"[\[\]:*?/\\]", "_", name)[:28] or "Curva"
    t, k = base, 2
    while t in used:
        t = f"{base[:25]}_{k}"
        k += 1
    used.add(t)
    return t


def to_xlsx(b: Bundle) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = "Datos"
    cols = PX_COLUMNS if b.pixels_only else COLUMNS
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in b.rows:
        ws.append([r.get(k) for k in cols])
    used = {"Datos", "Metadatos"}
    for e in b.entries:
        sh = wb.create_sheet(_sheet_title(e.curve_id, used))
        rows = [r for r in b.rows if r["curve_id"] == e.curve_id]
        if b.pixels_only:
            sh.append(["point", "px", "py", "origin"])
            for r in rows:
                sh.append([r["point"], r["px"], r["py"], r["origin"]])
        else:
            sh.append(["point", f"x [{rows[0]['x_unit']}]", f"y [{rows[0]['y_unit']}]", "origin"])
            for r in rows:
                sh.append([r["point"], r["x"], r["y"], r["origin"]])
        for c in sh[1]:
            c.font = Font(bold=True)
    meta = wb.create_sheet("Metadatos")
    img = b.image
    for k, v in [("Esquema", f"{SCHEMA} {SCHEMA_VERSION}"), ("Aplicación", f"PEMFC ModelLab {APP_VERSION}"),
                 ("Motor de digitalización", f"{ALGORITHM_ID} v{DIGITIZER_VERSION}"), ("Creado (UTC)", b.created_utc),
                 ("Imagen", img.filename), ("SHA-256 imagen", img.sha256), ("Tamaño [px]", f"{img.width}×{img.height}"),
                 ("Estado", "SIN CALIBRAR (solo píxeles)" if b.pixels_only else "Calibrado y confirmado"),
                 ("Aviso", "Datos digitalizados de una imagen; no constituyen una medición independiente.")]:
        meta.append([k, v])
    for d in {id(e.draft): e.draft for e in b.entries}.values():
        pm = _panel_meta(d)
        meta.append([])
        meta.append([f"Panel {pm['panel']}"])
        meta.append(["Unidades", f"x: {pm['units']['x']} · y: {pm['units']['y']} ({pm['units']['source']})"])
        cal = pm["calibration"]
        if cal:
            meta.append(["Calibración X", cal["x"]["formula"]])
            meta.append(["Calibración Y", cal["y"]["formula"]])
        for ax, rec in pm["axis_recognition"].items():
            meta.append([f"Reconocimiento {ax}", f"{rec['status']} · escala {rec['scale']} · {rec['n_inliers']} etiquetas"
                         + (f" · RMS {rec['rms_px']:.2f} px" if rec["rms_px"] is not None else "")])
        meta.append(["Ejes confirmados", json.dumps(pm["axis_confirmed"])])
    meta.column_dimensions["A"].width = 26
    meta.column_dimensions["B"].width = 90
    buf = io.BytesIO()
    with _exact_xlsx_numbers():
        wb.save(buf)
    return buf.getvalue()


class _exact_xlsx_numbers:
    """openpyxl escribe los floats con «%.16g» (16 cifras): no basta para la ida y vuelta exacta de un
    double (hasta 17). Durante el guardado se usa ``repr(float)``, válido en XLSX y exacto."""

    def __enter__(self):
        try:
            import openpyxl.cell._writer as w
        except ImportError:
            self.w = None
            return self
        self.w, self.orig = w, getattr(w, "safe_string", None)
        if self.orig is not None:
            def exact(value, _orig=self.orig):
                if isinstance(value, float) and math.isfinite(value):
                    return repr(value)
                return _orig(value)
            w.safe_string = exact
        return self

    def __exit__(self, *exc):
        if self.w is not None and self.orig is not None:
            self.w.safe_string = self.orig
        return False


def to_png(b: Bundle, max_side: int = 1800) -> bytes:
    img = Image.fromarray(b.image.rgb).convert("RGB")
    s = min(1.0, max_side / max(img.size))
    if s < 1:
        img = img.resize((round(img.width * s), round(img.height * s)), Image.Resampling.LANCZOS)
    dr = ImageDraw.Draw(img)
    for k, e in enumerate(b.entries):
        col = PALETTE[k % len(PALETTE)]
        pts = [(p["px"] * s, p["py"] * s) for p in e.points]
        if len(pts) >= 2:
            dr.line(pts, fill=col, width=2)
        for x, y in pts:
            dr.ellipse([x - 3.5, y - 3.5, x + 3.5, y + 3.5], fill=col, outline=(0, 0, 0))
        dr.rectangle([8, 8 + 18 * k, 22, 20 + 18 * k], fill=col)
        dr.text((28, 8 + 18 * k), e.curve_id + (" (SIN CALIBRAR)" if b.pixels_only else ""), fill=(0, 0, 0))
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def to_zip(b: Bundle) -> bytes:
    img = b.image
    buf = io.BytesIO()
    readme = (f"Exportación de PEMFC ModelLab {APP_VERSION} — digitalizador autónomo\n"
              f"Imagen: {img.filename} (SHA-256 {img.sha256})\nCurvas: {', '.join(e.curve_id for e in b.entries)}\n"
              f"Estado: {'SIN CALIBRAR (solo píxeles)' if b.pixels_only else 'calibrado y confirmado'}\n\n"
              "datos.csv / datos.xlsx / datos.json contienen los mismos valores. overlay.png superpone los puntos "
              "exportados a la imagen original. auditoria/ guarda el registro completo de cada panel (anclas, "
              "reconocimiento de ejes, ediciones manuales con marca de tiempo).\n"
              "Los datos digitalizados de una imagen no constituyen una medición independiente.\n")
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("LEEME.txt", readme)
        z.writestr("datos.csv", to_csv(b))
        z.writestr("datos.xlsx", to_xlsx(b))
        z.writestr("datos.json", to_json(b))
        z.writestr("overlay.png", to_png(b))
        z.writestr(f"original/{img.filename}", img.original_bytes)
        for d in {id(e.draft): e.draft for e in b.entries}.values():
            z.writestr(f"auditoria/panel_{_panel_label(d)}.json",
                       json.dumps(d.to_audit_dict(), ensure_ascii=False, indent=1))
    return buf.getvalue()


FORMATS = {"CSV": ("datos.csv", "text/csv", to_csv), "XLSX": ("datos.xlsx",
           "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", to_xlsx),
           "JSON": ("datos.json", "application/json", to_json), "PNG": ("overlay.png", "image/png", to_png),
           "ZIP": ("digitalizacion.zip", "application/zip", to_zip)}


def finite(v) -> bool:
    return isinstance(v, (int, float)) and math.isfinite(v)
